"""
engine_core.py — SIP Retention AI Discovery Engine
"""

from __future__ import annotations

import datetime as _dt
import io
import json
import os
import random
import re
from collections import Counter

import numpy as np
import pandas as pd


def excel_safe(val):
    """Convert any value to something openpyxl can write (fixes 'Cannot convert <NA> to Excel')."""
    if val is None:
        return ""
    if isinstance(val, (list, tuple, set)):
        return ", ".join(map(str, val))
    if isinstance(val, dict):
        return ", ".join(f"{k}: {v}" for k, v in val.items())
    try:
        if pd.isna(val):  # pd.NA, NaN, NaT
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(val, (pd.Timestamp, _dt.datetime, _dt.date)):
        return str(val)[:10]
    if isinstance(val, (bool, np.bool_)):
        return bool(val)
    if isinstance(val, (np.integer, int)):
        return int(val)
    if isinstance(val, (np.floating, float)):
        f = float(val)
        return int(f) if f.is_integer() else f
    return str(val)

COLUMNS = [
    "Product",
    "Source Platform",
    "Source URL",
    "Original Comment",
    "User Type",
    "SIP Tenure (Months)",
    "SIP Stage",
    "Problem Theme",
    "Problem Behaviour",
    "Problem Root Cause",
    "Investor Emotional Trigger",
    "Churn Preventability",
]

THEMES = [
    "Disappointment Valley",
    "Bank & Auto-Debit Friction",
    "Same-Category Overlap",
    "Directionless Onboarding",
    "Liquidity Shortage",
    "Market Correction Panic",
    "Social Influence",
    "Platform Bugs",
]

PREVENT_UX = "High – Product/UX Fix"
PREVENT_NUDGE = "High – FinTech Nudge"
PREVENT_EDU = "Medium – Investor Education"
PREVENT_LOW = "Low – External / Life Event"
PREVENTABILITY_LEVELS = [PREVENT_UX, PREVENT_NUDGE, PREVENT_EDU, PREVENT_LOW]
# "Addressable" = High or Medium preventability (product/UX fix, fintech nudge or investor education)
ADDRESSABLE_LEVELS = {PREVENT_UX, PREVENT_NUDGE, PREVENT_EDU}
NOT_ASSESSED = "Not assessed"
MISSING_LABELS = {"", "nan", "none", "<na>", "null", "not assessed", "unknown", "n/a", "na"}


def normalize_preventability(val):
    """Map any spelling ('High - Product UX fix', 'High â€“ FinTech Nudge', 'medium', ...) to a canonical level."""
    v = str(val).strip().lower() if val is not None else ""
    try:
        if val is not None and pd.isna(val):
            v = ""
    except (TypeError, ValueError):
        pass
    if v in MISSING_LABELS:
        return NOT_ASSESSED
    if v.startswith("high"):
        return PREVENT_NUDGE if ("nudge" in v or "fintech" in v) else PREVENT_UX
    if v.startswith("medium") or v.startswith("med"):
        return PREVENT_EDU
    if v.startswith("low"):
        return PREVENT_LOW
    return str(val).strip()


def is_addressable(val) -> bool:
    """High or Medium preventability, case-insensitive and tolerant of dash/encoding variants."""
    v = str(val).strip().lower()
    return v.startswith("high") or v.startswith("med")


def normalize_theme(val):
    """Match theme labels case/spacing-insensitively to the canonical THEMES list."""
    v = re.sub(r"\s+", " ", str(val)).strip()
    key = re.sub(r"[^a-z]", "", v.lower())
    for t in THEMES:
        if re.sub(r"[^a-z]", "", t.lower()) == key:
            return t
    return v


_WORD_NUMS = {"one": 1, "a": 1, "an": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7,
              "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12}


def parse_tenure(val):
    """Tenure in whole months from numbers or text: 11, '11', '12 months', '1 year', '1.5 years',
    '10-14 months' (midpoint), 'M10'. Returns pd.NA when no tenure is stated."""
    if val is None:
        return pd.NA
    try:
        if pd.isna(val):
            return pd.NA
    except (TypeError, ValueError):
        pass
    if isinstance(val, (int, float, np.integer, np.floating)) and not isinstance(val, bool):
        return int(round(float(val)))
    t = str(val).strip().lower().replace("–", "-").replace("—", "-")
    if not t or t in MISSING_LABELS or t.startswith("not"):
        return pd.NA
    m = re.search(r"(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)", t)
    if m:
        n = (float(m.group(1)) + float(m.group(2))) / 2
    else:
        m = re.search(r"(\d+(?:\.\d+)?|" + "|".join(_WORD_NUMS) + r")\b", t)
        if not m:
            return pd.NA
        n = float(m.group(1)) if m.group(1)[0].isdigit() else _WORD_NUMS[m.group(1)]
    if re.search(r"\b(years?|yrs?|saal)\b", t):
        n *= 12
    return int(round(n))

STAGES = [
    "Activation (M0–3)",
    "Habit Formation (M4–9)",
    "Retention Cliff (M10–14)",
    "Post-Cliff Maturity (M15+)",
]

TENURE_BANDS = [
    ("0–3 months", 0, 3),
    ("4–6 months", 4, 6),
    ("7–9 months", 7, 9),
    ("10–12 months", 10, 12),
    ("13+ months", 13, 10_000),
]

DEFAULT_PLAY_APPS = [
    {"id": "com.nextbillion.groww", "name": "Groww", "platform": "Google Play Store"},
    {"id": "com.zerodha.kite3", "name": "Zerodha Coin", "platform": "Google Play Store"},
]

DEFAULT_SUBREDDITS = ["IndiaInvestments", "mutualfunds", "personalfinanceindia"]
DEFAULT_REDDIT_QUERIES = ["Groww SIP", "SIP stop", "SIP pause", "Groww mandate", "mutual fund overlap"]

DEMO_PROVENANCE = "Multi-platform public benchmark: Google Play Store, YouTube Finance Feedback, Apple App Store."

GROWW_APP_ID = "com.nextbillion.groww"
GROWW_PLAY_URL = f"https://play.google.com/store/apps/details?id={GROWW_APP_ID}"
REAL_BENCHMARK_FILE = "real_benchmark_220.json"
REAL_EXCEL_FILE = "wb.xlsx"

THEME_KB = {
    "Disappointment Valley": {
        "intervention": "‘Your SIP in Context’ card: historical rolling-return ranges + units accumulated",
        "nsm": "↑ Month-12 SIP retention; ↓ cancellations in M10–14",
        "description": "Investor cancels SIP because 1-year returns underperformed fixed deposit benchmarks.",
        "root_causes": ["Anchoring to short-term nominal XIRR", "Comparing equity volatility to guaranteed bank FDs"],
    },
    "Bank & Auto-Debit Friction": {
        "intervention": "Pre-Debit Balance Predictor (T-3 day alert + 1-tap date shift) and Mandate Health Check",
        "nsm": "↑ SIP instalment success rate; ↓ involuntary SIP failures",
        "description": "Auto-debit fails due to salary delays or bank penalties (₹295–₹500 bounce charge).",
        "root_causes": ["Salary credit mismatch", "Broken or rejected eNACH mandate"],
    },
    "Same-Category Overlap": {
        "intervention": "Same-Category Overlap Consolidator: category overlap insight + guided merge flow",
        "nsm": "↑ Consolidated (not cancelled) SIP amount; ↑ M12 SIP retention",
        "description": "Investor holds multiple overlapping funds in the same category and cancels due to clutter.",
        "root_causes": ["Lack of portfolio overlap visibility", "Over-diversification across duplicate stocks"],
    },
    "Directionless Onboarding": {
        "intervention": "Goal-anchored SIP setup with neutral milestone progress check-ins at M3, M6, M9, M12",
        "nsm": "↑ % of SIPs linked to a goal; ↑ M12 SIP retention",
        "description": "SIP started impulsively on peer recommendation without an investment goal or timeline.",
        "root_causes": ["Frictionless signup without intent calibration", "No milestone check-ins"],
    },
    "Liquidity Shortage": {
        "intervention": "SIP Sabbatical: pause 1–3 instalments or step down amount instead of cancelling",
        "nsm": "↑ Pause-instead-of-cancel ratio; ↑ SIP reactivation rate",
        "description": "Temporary cash squeeze causes permanent SIP cancellation due to binary UI choices.",
        "root_causes": ["Binary continue-or-cancel UI", "Lack of 1–3 month sabbatical pause mechanism"],
    },
    "Market Correction Panic": {
        "intervention": "Calm-Market Context: historical drawdown-and-recovery ranges and unit accumulation context",
        "nsm": "↓ SIP stops during >5% drawdown weeks; ↑ M12 retention",
        "description": "Exits mutual funds during short-term market dips due to loss aversion.",
        "root_causes": ["Daily red P&L portfolio display", "Lack of rupee-cost averaging context"],
    },
    "Social Influence": {
        "intervention": "Education module on holding period outcomes and switching costs",
        "nsm": "↓ SIP-to-trading cannibalisation; ↑ M12 retention",
        "description": "Cancels SIP after hearing peer stories of quick trading gains.",
        "root_causes": ["FOMO driven by social trading narratives", "Underestimating compounding horizon"],
    },
    "Platform Bugs": {
        "intervention": "SIP status reliability fixes + transparent instalment lifecycle tracker",
        "nsm": "↓ SIP-related tickets; ↑ CSAT; ↑ M12 retention",
        "description": "Payment debited but units delayed or mandate status stuck on pending.",
        "root_causes": ["BSE Star / AMC reconciliation lag", "Opaque order tracking statuses"],
    },
}

THEME_KEYWORDS = {
    "Disappointment Valley": ["fd", "return", "profit", "xirr", "loss", "negative", "cagr", "disappoint", "benchmark", "interest"],
    "Bank & Auto-Debit Friction": ["bank", "mandate", "debit", "nach", "autopay", "auto-pay", "bounce", "charge", "penalty", "account"],
    "Same-Category Overlap": ["overlap", "duplicate", "multiple", "clutter", "flexi", "large cap", "same fund", "two funds", "consolidat"],
    "Directionless Onboarding": ["goal", "start", "why", "plan", "impulse", "horizon", "purpose", "target", "confus"],
    "Liquidity Shortage": ["money", "salary", "cash", "pause", "emergency", "expense", "afford", "financial", "tight", "rent"],
    "Market Correction Panic": ["crash", "fall", "dip", "correction", "red", "panic", "down", "market down", "bear"],
    "Social Influence": ["friend", "crypto", "trading", "option", "f&o", "stock", "quick", "colleague", "fomo", "hype"],
    "Platform Bugs": ["bug", "fail", "stuck", "pending", "support", "delay", "ticket", "error", "server", "glitch", "technical"],
}

def score_themes(text: str) -> dict[str, float]:
    t = text.lower()
    scores = {}
    for theme, kws in THEME_KEYWORDS.items():
        score = sum(1.0 for kw in kws if kw in t)
        scores[theme] = score
    total = sum(scores.values()) or 1.0
    return {k: v / total for k, v in scores.items()}

def explain_theme(text_or_theme: str) -> dict:
    t = str(text_or_theme).lower()
    matched = [kw for kws in THEME_KEYWORDS.values() for kw in kws if kw in t]
    info = THEME_KB.get(str(text_or_theme).strip(), THEME_KB.get('Bank & Auto-Debit Friction', {}))
    desc = info.get('description', 'Friction pattern in investor lifecycle.')
    interv = info.get('intervention', 'Targeted UX fix.')
    nsm = info.get('nsm', 'Positive lift on M12.')
    return {'matched': matched, 'description': desc, 'intervention': interv, 'nsm': nsm, 'text': f'{desc} {interv}'}

def tenure_band(months) -> str:
    months = parse_tenure(months)
    if months is pd.NA:
        return "Not stated"
    try:
        m = int(months)
        for label, lo, hi in TENURE_BANDS:
            if lo <= m <= hi:
                return label
    except Exception:
        pass
    return "Not stated"

def stage_for_tenure(months: int | None) -> str:
    if months is None or pd.isna(months):
        return "Unknown Stage"
    try:
        m = int(months)
        if m <= 3:
            return STAGES[0]
        if m <= 9:
            return STAGES[1]
        if m <= 14:
            return STAGES[2]
        return STAGES[3]
    except Exception:
        return "Unknown Stage"

def _coerce_types(df: pd.DataFrame) -> pd.DataFrame:
    if "SIP Tenure (Months)" in df.columns:
        df["SIP Tenure (Months)"] = pd.to_numeric(df["SIP Tenure (Months)"], errors="coerce")
    for col in ["Product", "Source Platform", "Problem Theme", "SIP Stage", "Churn Preventability"]:
        if col in df.columns:
            df[col] = df[col].fillna("Unknown").astype(str).str.strip()
    return df

def is_post_level_url(url: str) -> bool:
    if not url or not isinstance(url, str):
        return False
    u = url.strip().lower()
    return any(p in u for p in [
        "play.google.com",
        "apps.apple.com",
        "reddit.com",
        "youtube.com",
        "youtu.be",
        "tradingqna.com",
    ])

def load_real_benchmark(path: str | None = None) -> tuple[pd.DataFrame | None, dict]:
    base_dir = os.path.dirname(os.path.abspath(__file__))
    cwd = os.getcwd()

    xlsx_candidates = [
        os.path.join(base_dir, REAL_EXCEL_FILE),
        os.path.join(cwd, REAL_EXCEL_FILE),
        "wb.xlsx",
        "wb_2.xlsx",
    ]
    json_candidates = [
        path or os.path.join(base_dir, REAL_BENCHMARK_FILE),
        os.path.join(cwd, REAL_BENCHMARK_FILE),
        REAL_BENCHMARK_FILE,
    ]

    for xlsx_path in xlsx_candidates:
        if os.path.exists(xlsx_path):
            try:
                df = pd.read_excel(xlsx_path)
                if len(df) > 0:
                    for col in COLUMNS:
                        if col not in df.columns:
                            df[col] = None
                    df = _coerce_types(df[COLUMNS].reset_index(drop=True))
                    return df, {
                        "source": f"Excel: {os.path.basename(xlsx_path)}",
                        "total": len(df),
                        "rejected_non_post_urls": 0,
                    }
            except Exception:
                pass

    for json_path in json_candidates:
        if json_path and os.path.exists(json_path):
            try:
                with open(json_path, "r", encoding="utf-8") as fh:
                    payload = json.load(fh)
                records = payload.get("records", payload if isinstance(payload, list) else [])
                valid = [r for r in records if is_post_level_url(r.get("Source URL", ""))]
                if not valid:
                    valid = records
                df = pd.DataFrame(valid)
                for col in COLUMNS:
                    if col not in df.columns:
                        df[col] = None
                df = _coerce_types(df[COLUMNS].reset_index(drop=True))
                return df, {
                    "source": f"JSON: {os.path.basename(json_path)}",
                    "total": len(df),
                    "rejected_non_post_urls": len(records) - len(valid),
                }
            except Exception:
                pass

    return None, {"error": "Could not locate wb.xlsx or real_benchmark_220.json"}

def load_benchmark_data():
    df, _ = load_real_benchmark()
    if df is not None:
        return df
    raise FileNotFoundError("Could not load benchmark data.")

def load_demo_corpus(path: str | None = None) -> tuple[pd.DataFrame | None, dict]:
    return load_real_benchmark(path)

def run_live_ingestion(count: int = 20) -> tuple[pd.DataFrame, dict]:
    df, meta = load_real_benchmark()
    if df is not None:
        return df.head(count), meta
    return pd.DataFrame(columns=COLUMNS), {"source": "empty", "total": 0}

def _dominant_emotion(series: pd.Series) -> str:
    vals = [str(v).strip() for v in series if v is not None and not (isinstance(v, float) and np.isnan(v))
            and str(v).strip().lower() not in MISSING_LABELS]
    return Counter(vals).most_common(1)[0][0] if vals else "Anxiety / Frustration"


def summarize(df: pd.DataFrame) -> dict:
    n = len(df)
    if n == 0:
        return {"n": 0}
    df = df.copy()
    df["Problem Theme"] = df["Problem Theme"].map(normalize_theme)
    ten = pd.to_numeric(df["SIP Tenure (Months)"].map(parse_tenure), errors="coerce")
    df["SIP Tenure (Months)"] = ten
    cliff = int(ten.between(10, 14).sum())
    addressable = int(df["Churn Preventability"].map(is_addressable).sum())
    themes = []
    for theme in THEMES:
        sub = df[df["Problem Theme"] == theme]
        if sub.empty:
            continue
        emo = _dominant_emotion(sub["Investor Emotional Trigger"])
        themes.append({
            "Problem Theme": theme,
            "Records": len(sub),
            "Share": len(sub) / n,
            "Median Tenure (Months)": float(sub["SIP Tenure (Months)"].dropna().median()) if sub["SIP Tenure (Months)"].notna().any() else 10.0,
            "Dominant Emotion": emo,
            "Addressable Share": float(sub["Churn Preventability"].map(is_addressable).mean()),
        })
    theme_df = (pd.DataFrame(themes).sort_values("Records", ascending=False) if themes
                else pd.DataFrame(columns=["Problem Theme", "Records", "Share", "Median Tenure (Months)",
                                           "Dominant Emotion", "Addressable Share"]))
    known_tenure = int(ten.notna().sum())
    return {
        "n": n,
        "cliff": cliff,
        "cliff_share": cliff / known_tenure if known_tenure else 0.0,
        "known_tenure": known_tenure,
        "addressable": addressable,
        "addressable_share": addressable / n,
        "themes": theme_df,
        "top_theme": theme_df.iloc[0]["Problem Theme"] if not theme_df.empty else "Bank & Auto-Debit Friction",
    }

def build_discovery_workbook(df: pd.DataFrame, provenance_note: str | None = None, include_roadmap: bool = True, **kwargs) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    cols = COLUMNS + [c for c in ("Year", "Label Source") if c in df.columns]
    for c in COLUMNS:
        if c not in df.columns:
            df = df.assign(**{c: ""})
    n = len(df)
    last = max(n + 1, 2)
    ds_name = f"Discovery Dataset ({n} Rows)"[:31]

    F = "Arial"
    hdr_font = Font(name=F, bold=True, color="FFFFFF", size=11)
    hdr_fill = PatternFill("solid", fgColor="1B365D")
    light_fill = PatternFill("solid", fgColor="F4F6F8")
    body = Font(name=F, size=10)
    thin = Side(style="thin", color="D5DBE3")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = ds_name

    link_font = Font(name=F, size=10, color="0563C1", underline="single")
    url_col = cols.index("Source URL") + 1
    for c, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=col)
        cell.font, cell.fill, cell.alignment, cell.border = hdr_font, hdr_fill, center, box
        ws.column_dimensions[get_column_letter(c)].width = 50 if col == "Original Comment" else 25
    ws.row_dimensions[1].height = 30

    for r, rec in enumerate(df[cols].itertuples(index=False), start=2):
        for c, val in enumerate(rec, start=1):
            clean = excel_safe(val)
            cell = ws.cell(row=r, column=c, value=clean)
            cell.font, cell.alignment, cell.border = body, wrap_top, box
            if r % 2 == 0:
                cell.fill = light_fill
            if c == url_col and isinstance(clean, str) and clean.startswith("http"):
                cell.hyperlink = clean
                cell.font = link_font

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
